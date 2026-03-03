from __future__ import annotations

from datetime import date
from io import BytesIO

import pandas as pd
from docx import Document

from disclosures import build_disclosure_table, maturity_analysis


def create_note_docx(
    schedules: pd.DataFrame,
    fiscal_year_end: date,
    narrative: str,
    comparative_schedules: pd.DataFrame | None = None,
) -> bytes:
    doc = Document()
    doc.add_heading("(7)    LEASES", level=2)
    doc.add_paragraph(narrative)

    current = build_disclosure_table(schedules, fiscal_year_end)

    for section in ["Financing Leases", "Operating Leases"]:
        doc.add_heading(section, level=3)

    doc.add_paragraph("Lease cost table")
    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = str(fiscal_year_end.year)
    table.rows[0].cells[2].text = str(fiscal_year_end.year - 1)

    metrics = ["Amortization of ROU assets", "Interest on lease liabilities", "Operating lease expense"]
    prior = build_disclosure_table(comparative_schedules, date(fiscal_year_end.year - 1, 12, 31)) if comparative_schedules is not None else pd.DataFrame(columns=["Label", "Amount"])

    for m in metrics:
        row = table.add_row().cells
        row[0].text = m
        row[1].text = f"{float(current[current['Label'] == m]['Amount'].sum()):,.2f}"
        row[2].text = f"{float(prior[prior['Label'] == m]['Amount'].sum()):,.2f}"

    doc.add_paragraph("Weighted average remaining term and discount rates")
    wa = doc.add_table(rows=1, cols=2)
    wa.rows[0].cells[0].text = "Metric"
    wa.rows[0].cells[1].text = "Amount"
    for label in [
        "Weighted-average remaining lease term in years for finance leases",
        "Weighted-average remaining lease term in years for operating leases",
        "Weighted-average discount rate for finance leases",
        "Weighted-average discount rate for operating leases",
    ]:
        r = wa.add_row().cells
        r[0].text = label
        r[1].text = f"{float(current[current['Label'] == label]['Amount'].sum()):,.4f}"

    doc.add_paragraph("Maturity analysis")
    mat = maturity_analysis(schedules, fiscal_year_end)
    mt = doc.add_table(rows=1, cols=3)
    mt.rows[0].cells[0].text = "Period"
    mt.rows[0].cells[1].text = "Finance"
    mt.rows[0].cells[2].text = "Operating"
    for _, x in mat.iterrows():
        rr = mt.add_row().cells
        rr[0].text = str(x["Label"])
        rr[1].text = f"{x['Finance']:,.2f}"
        rr[2].text = f"{x['Operating']:,.2f}"

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()
