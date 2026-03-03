from __future__ import annotations

from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from disclosures import build_disclosure_table, maturity_analysis

DATA_COLUMNS = [
    "Date",
    "Month",
    "Amortization Expense (Finance) or Operating Lease Expense (Operating)",
    "ROU Asset Beginning",
    "ROU Asset",
    "ROU Asset EOM",
    "LT Liability Beginning",
    "Interest Expense (Not Applicable for Operating Lease)",
    "LT Liability (interest)",
    "LT Liability (Payment at BOM)",
    "LT Liability (Payment at EOM)",
    "Total ST & LT Liability EOM",
    "LT Liability",
    "ST Liability",
    "LT Liability EOM",
    "ST Liability EOM",
    "Cash/AP for Lease Payment at BOM",
    "Cash/AP for Lease Payments at EOM",
    "Cash/AP for Variable Lease Expense Payment",
    "Cash/AP for Non-Lease Payment",
    "Cash/AP (Summary of all Cash/AP Entries)",
    "Variable Lease Expense",
    "Other P&L Accounts",
    "Other BS Accounts",
]


def create_audit_workbook(
    schedules: pd.DataFrame,
    fiscal_year_end: date,
    narrative_lines: list[str],
) -> bytes:
    wb = Workbook()
    ws_data = wb.active
    ws_data.title = "Data"

    for idx, c in enumerate(DATA_COLUMNS, start=1):
        ws_data.cell(row=1, column=idx, value=c).font = Font(bold=True)

    fiscal_year = fiscal_year_end.year
    year_df = schedules[pd.to_datetime(schedules["date"]).dt.year == fiscal_year].copy()

    for i, (_, row) in enumerate(year_df.iterrows(), start=2):
        values = [
            row["date"],
            row["month"],
            row["lease_expense"],
            row["rou_begin"],
            row["rou_amortization"],
            row["rou_end"],
            row["begin_liability"],
            row["interest_expense"],
            row["interest_expense"],
            row["payment_bom"],
            row["payment_eom"],
            row["end_liability"],
            row["lt_liability"],
            row["st_liability"],
            row["lt_liability"],
            row["st_liability"],
            row["payment_bom"],
            row["payment_eom"],
            row["variable_payment"],
            row["nonlease_payment"],
            row["payment_bom"] + row["payment_eom"] + row["variable_payment"] + row["nonlease_payment"],
            row["variable_payment"],
            0,
            0,
        ]
        for j, v in enumerate(values, start=1):
            ws_data.cell(row=i, column=j, value=v)

    line_start = len(year_df) + 4
    for i, text in enumerate(narrative_lines):
        ws_data.cell(row=line_start + i, column=1, value=text)

    ws_data.freeze_panes = "A2"

    ws_disc = wb.create_sheet(title=f"Disclosures - {fiscal_year}")
    ws_disc["A1"] = "FASB ASC 842 Footnote"
    ws_disc["A1"].font = Font(bold=True)
    ws_disc["A2"] = "Year Ending"
    ws_disc["B2"] = f"{fiscal_year}-12"

    disc = build_disclosure_table(schedules, fiscal_year_end)
    r = 4
    for _, row in disc.iterrows():
        ws_disc.cell(row=r, column=1, value=row["Label"])
        ws_disc.cell(row=r, column=2, value=row["Amount"])
        r += 1

    ws_disc.cell(row=r + 1, column=1, value="Maturity Analysis")
    ws_disc.cell(row=r + 1, column=2, value="Finance")
    ws_disc.cell(row=r + 1, column=3, value="Operating")
    m = maturity_analysis(schedules, fiscal_year_end)
    for i, (_, row) in enumerate(m.iterrows(), start=r + 2):
        ws_disc.cell(row=i, column=1, value=row["Label"])
        ws_disc.cell(row=i, column=2, value=row["Finance"])
        ws_disc.cell(row=i, column=3, value=row["Operating"])

    ws_disc.cell(row=r + 2 + len(m) + 1, column=1, value="* Not calculated by LeaseCrunch")
    ws_disc.freeze_panes = "A4"

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
